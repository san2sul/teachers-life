src=open('audit2.py').read();exec(src.split("sizes=")[0])
import openpyxl,io,csv
S["list"][0]["data"]["obs"]=[{"id":"o1","sid":"10101","date":"2026-10-15","tags":{},"memo":"발표 적극적"}]
S["list"][0]["data"]["sessions"]=[{"id":"s1","date":"2026-10-16","time":"09:00","p":{"10101":1,"10102":0},"absent":["10102"]}]
def st(T):return json.loads(T.evaluate("localStorage.getItem('att-app-v2')"))
D=SP+"/dl";os.makedirs(D,exist_ok=True)
def grab(T,fn):
  with T.expect_download(timeout=20000) as d:fn()
  p=D+"/"+d.value.suggested_filename;d.value.save_as(p);return p
def show_x(p,n=5):
  wb=openpyxl.load_workbook(p)
  for ws in wb.worksheets:
    print("  [%s] merges=%s"%(ws.title,[str(m) for m in ws.merged_cells.ranges][:4]))
    for r in ws.iter_rows(min_row=1,max_row=n,values_only=True):print("   ",[x for x in r if x not in (None,"")][:8])
    print("    A2 align:",ws["A2"].alignment.horizontal)
with sync_playwright() as p:
  b=p.chromium.launch();c=b.new_context(viewport={"width":390,"height":844},service_workers="block",accept_downloads=True,permissions=["clipboard-read","clipboard-write"]);c.route("**/*",handler);T=c.new_page();T.on("pageerror",lambda e:errs.append(str(e)))
  T.on("dialog",lambda d:d.accept("임시이름"))
  T.goto("https://tl.test/index.html");T.evaluate("s=>{localStorage.setItem('att-app-v2',s);localStorage.setItem('att-app-lastver','9.9.9')}",json.dumps(S));T.reload();T.wait_for_timeout(400)
  # 1) 이름 없이 받으면 prompt
  T.click(".navb[data-v=set]");T.evaluate("document.getElementById('d-io').open=true") if T.query_selector("#d-io") else None
  p1=grab(T,lambda:T.evaluate("document.querySelector('[data-tpl=stu][data-f=xlsx]').click()"));print("stu xlsx:",os.path.basename(p1));show_x(p1)
  print("me after prompt:",T.evaluate("localStorage.getItem('att-app-me')"))
  T.fill("#meName","홍길동");T.wait_for_timeout(700);print("meEx:",T.inner_text("#meEx"),"| tag:",T.inner_text("#meTag"))
  p2=grab(T,lambda:T.evaluate("document.querySelector('[data-tpl=staff][data-f=csv]').click()"));print("staff csv:");print(open(p2,encoding='utf-8-sig').read()[:300])
  # re-upload stu template filled
  wb=openpyxl.load_workbook(p1);ws=wb.worksheets[0];ws.append(["10109","신입생","010-5555-6666","new@ex.com"]);ws["A5"]="10199";ws["B5"]="홍길동(예시)";wb.save(D+"/stu_filled.xlsx")
  T.evaluate("document.querySelector('[data-bulk=stu]').click()");T.wait_for_timeout(300)
  T.set_input_files("#bkFile",D+"/stu_filled.xlsx");T.wait_for_timeout(800);print("bulk:",T.inner_text("#sheetBox")[:300].replace("\n"," | "))
  T.evaluate("document.getElementById('sheet').hidden=true")
  # eval template by class
  T.click(".navb[data-v=stat]");T.click("#stabs .chip[data-t=eval]");T.click("#evTpl");T.wait_for_timeout(200)
  T.evaluate("document.querySelector('#tpMode [data-m=cls]')&&document.querySelector('#tpMode [data-m=cls]').click()")
  p3=grab(T,lambda:T.click("#tpOk"));print("eval tpl:");show_x(p3,6)
  wb=openpyxl.load_workbook(p3);ws=wb["수행평가"];print("  formula:",[ws.cell(5,c).value for c in range(1,9)])
  for r in range(5,5+8):
    for cc in range(4,7):ws.cell(r,cc).value=10
  wb.save(D+"/ev_filled.xlsx")
  T.set_input_files("#evFile",D+"/ev_filled.xlsx");T.wait_for_timeout(1500);print("UPLOAD:",T.inner_text("#sheetBox")[:300].replace("\n"," | "))
  T.fill("#evReply","t@gmail.com");T.click("#evOk");T.wait_for_timeout(500)
  ev=[e for e in st(T)['evals'] if e.get('file',{}).get('name','').startswith('ev_filled')][0];print("  parsed stu:",len(ev['stu']),ev['stu'][0]['vals'],"hdr",ev['file'].get('hdr'))
  # export signed xlsx (our template)
  def signed(label):
    T.click("#evXl");T.wait_for_timeout(300);p4=grab(T,lambda:T.click("#eoOk"));print("signed",label);show_x(p4,7)
    wb=openpyxl.load_workbook(p4);ws=wb.worksheets[0];print("   formulas:",[c.value for row in ws.iter_rows() for c in row if isinstance(c.value,str) and c.value.startswith('=')][:4])
  signed("ours")
  for f in ["ev_school.xlsx","eval_test.xlsx"]:
    T.click("#evBack");T.set_input_files("#evFile",SP+"/"+f);T.wait_for_timeout(1500);print("UP",f,T.inner_text("#sheetBox")[:90].replace("\n"," "))
    T.fill("#evReply","t@gmail.com");T.click("#evOk");T.wait_for_timeout(500);signed(f)
    wb=openpyxl.load_workbook(SP+"/"+f);ws=wb.worksheets[0];print("   ORIGINAL merges:",[str(m) for m in ws.merged_cells.ranges][:6])
  # sign summary
  T.click("#evBack");T.click("#evSum");p5=grab(T,lambda:T.click("#smX"));print("summary:");show_x(p5,4)
  T.evaluate("document.getElementById('sheet').hidden=true")
  # obs
  T.click(".navb[data-v=att]");T.click("#amode [data-m=obs]");T.wait_for_timeout(200);p6=grab(T,lambda:T.click("#obsXl"));print("obs:");show_x(p6,4)
  # stats csv
  T.click(".navb[data-v=stat]");T.click("#stabs .chip[data-t=att]");T.wait_for_timeout(200)
  T.evaluate("document.querySelectorAll('#s-att [data-r]').forEach(b=>{if(b.dataset.r==='all')b.click()})");T.wait_for_timeout(200)
  try:
    p7=grab(T,lambda:T.click("#sendFile"));print("stats csv:");print(open(p7,encoding='utf-8-sig').read()[:250])
  except Exception as e:print("stats err",e,T.inner_text("#toast") if T.query_selector('#toast') else '')
  print("report:",T.evaluate("1")) 
  b.close()
print(errs)
