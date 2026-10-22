src=open('audit2.py').read();exec(src.split("sizes=")[0])
import openpyxl,datetime
tt=S["tt"]["list"][0];tt["chg"]=[];tt["start"]="09:00";tt["classMin"]=45;tt["breakMin"]=10;tt["lunchMin"]=60
tt["cells"]={"0-2":{"subject":"물리","clsId":"c1","clsName":"","room":""},"4-2":{"subject":"통합과학","clsId":"c2","clsName":"","room":""},"4-5":{"subject":"물리","clsId":"c1","clsName":"","room":""},"4-6":{"subject":"물리","clsId":"c1","clsName":"","room":""},"0-1":{"subject":"과학탐구","clsId":"c2","clsName":"","room":""}}
D=lambda y,m,d:datetime.datetime(y,m,d)
with sync_playwright() as p:
  b=p.chromium.launch();c=b.new_context(viewport={"width":390,"height":844},device_scale_factor=2,service_workers="block",accept_downloads=True);c.route("**/*",handler);T=c.new_page();T.on("pageerror",lambda e:errs.append(str(e)))
  T.goto("https://tl.test/index.html");T.evaluate("s=>{localStorage.setItem('att-app-v2',s);localStorage.setItem('att-app-lastver','9.9.9');localStorage.setItem('att-app-me','김물리');localStorage.setItem('att-app-setup','1')}",json.dumps(S));T.reload();T.wait_for_timeout(500)
  T.click(".navb[data-v=tt]");T.click("#acUp2")
  with T.expect_download() as dl:T.click("#acTpl")
  fp=SP+"/tpl2.xlsx";dl.value.save_as(fp);wb=openpyxl.load_workbook(fp);ws=wb["학사일정"];print("sheets:",wb.sheetnames)
  for r in ws.iter_rows(min_row=4,max_row=11,values_only=True):print("  ",r)
  print("dv:",sorted(set(str(dv.formula1) for dv in ws.data_validations.dataValidation)))
  for r in range(5,12):
    for cc in range(1,11):ws.cell(r,cc).value=None
  rows=[(D(2026,11,6),None,"금요일 2교시 월요일 수업으로 운영","요일 교체","2","월요일",None,None,None,None),
        (D(2026,10,23),D(2026,12,18),"동아리 활동","교시 행사","5,6",None,"매주",None,None,None),
        (D(2026,11,13),None,"창의적 체험활동","교시 행사","5-6",None,None,"1학년 1반",None,None),
        (D(2026,11,20),None,"체육대회 (오후)","수업 없음","5-7",None,None,None,None,None),
        (D(2026,11,5),None,"학부모 상담","일정만","7",None,None,None,None,None),
        (D(2026,11,9),None,"월요일을 금요일 시간표로 운영","요일 교체",None,"금요일",None,None,None,None),
        (D(2026,12,24),None,"단축수업","단축수업",None,None,None,None,35,None)]
  for i,row in enumerate(rows):
    for j,v in enumerate(row):ws.cell(5+i,1+j).value=v
  wb.save(SP+"/tpl2_filled.xlsx")
  T.set_input_files("#acIn",SP+"/tpl2_filled.xlsx");T.wait_for_timeout(800)
  for x in T.evaluate("[...document.querySelectorAll('#acList .acrow')].map(r=>r.querySelector('.acd').innerText.replace(/\\n/g,'')+' | '+r.querySelector('[data-t]').value+' | '+((r.querySelector('.acmin')||{}).innerText||'')+' | '+r.querySelector('select').selectedOptions[0].text)"):print("   ",x)
  print("sum:",T.inner_text("#acSum"));T.screenshot(path=SP+"/shot/ac2_review.png")
  T.click("#acOk");T.wait_for_timeout(300);print("toast:",T.evaluate("(document.querySelector('.toast')||{}).textContent"))
  st=T.evaluate("JSON.parse(localStorage.getItem('att-app-v2'))");t=st['tt']['list'][0]
  for x in t['chg']:print("  chg",x['scope'],x['d'],x['p'],x['from'],x['to'],x['type'],x.get('subject'),x.get('clsId'),x.get('act',''),x.get('sw',''),x['reason'])
  print("  events:",[(e['date'],e['start'],e['end'],e['title']) for e in st['events'] if e['id'].startswith('ac')])
  print("  acal:",len(t.get('acal') or st['tt'].get('acal') or []))
  # timetable week of 11/2
  T.evaluate("()=>{}");
  for k in range(2):T.click("#ttNext");T.wait_for_timeout(80)
  print("week:",T.inner_text("#ttWeekTxt"));T.screenshot(path=SP+"/shot/ac2_week.png",full_page=True)
  # roundtrip
  T.click("#acUp2")
  with T.expect_download() as dl:T.click("#acTpl")
  fp=SP+"/tpl2_round.xlsx";dl.value.save_as(fp);ws=openpyxl.load_workbook(fp)["학사일정"]
  for r in ws.iter_rows(min_row=12,max_row=19,values_only=True):print("  R",r)
  T.set_input_files("#acIn",fp);T.wait_for_timeout(800);T.click("#acOk");T.wait_for_timeout(300)
  st=T.evaluate("JSON.parse(localStorage.getItem('att-app-v2'))");t=st['tt']['list'][0];print("after round chg:",len(t['chg']),"acal:",len(st['tt']['acal']))
  # bundle to receiver with different timetable
  T.click(".navb[data-v=set]");T.evaluate("document.getElementById('d-share').open=true")
  T.evaluate("()=>{for(const id of ['shStu','shTc','shStaff','shTags']){const e=document.getElementById(id);if(e.checked)e.click()} }")
  if not T.is_checked("#shCal"):T.click("#shCal")
  T.click("#mkLink");T.wait_for_timeout(400);url=T.inner_text("#linkOut");open(SP+"/bundle_url.txt","w").write(url)
  R=json.loads(json.dumps(S));R["tt"]["list"][0]["cells"]={"0-2":{"subject":"국어","clsId":"c1","clsName":"","room":""},"4-2":{"subject":"수학","clsId":"c1","clsName":"","room":""}};R["tt"]["list"][0]["chg"]=[]
  c2=b.new_context(viewport={"width":390,"height":844},service_workers="block");c2.route("**/*",handler);B=c2.new_page();B.on("pageerror",lambda e:errs.append("B:"+str(e)))
  B.goto("https://tl.test/index.html");B.evaluate("s=>{localStorage.setItem('att-app-v2',s);localStorage.setItem('att-app-setup','1');localStorage.setItem('att-app-me','이국어')}",json.dumps(R));B.goto('about:blank');B.goto(url);B.wait_for_timeout(1000)
  print("recv:",B.inner_text("#sheetBox")[:200].replace("\n"," | "));B.click("#rcOk");B.wait_for_timeout(300)
  st=B.evaluate("JSON.parse(localStorage.getItem('att-app-v2'))");t=st['tt']['list'][0]
  print("recv chg:",[(x["d"],x["p"],x["from"],x["type"],x.get("subject"),x.get("sw")) for x in t["chg"]]);print("errbar:",B.evaluate("(()=>{const e=document.getElementById('errBar');return e&&!e.hidden?e.innerText:''})()"));print("recv cells:",t["cells"],"acal:",[(a["title"],a["kind"]) for a in (st["tt"].get("acal") or [])])
  b.close()
print(errs)
