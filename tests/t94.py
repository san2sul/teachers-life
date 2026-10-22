src=open('audit2.py').read();exec(src.split("sizes=")[0])
import openpyxl,datetime
with sync_playwright() as p:
  b=p.chromium.launch();c=b.new_context(viewport={"width":390,"height":844},device_scale_factor=2,service_workers="block",accept_downloads=True);c.route("**/*",handler);T=c.new_page();T.on("pageerror",lambda e:errs.append(str(e)))
  T.goto("https://tl.test/index.html");T.evaluate("s=>{localStorage.setItem('att-app-v2',s);localStorage.setItem('att-app-lastver','9.9.9');localStorage.setItem('att-app-me','김물리');localStorage.setItem('att-app-setup','1')}",json.dumps(S));T.reload();T.wait_for_timeout(500)
  T.click(".navb[data-v=tt]");T.click("#acUp2");T.wait_for_timeout(100);T.screenshot(path=SP+"/shot/ac_open.png")
  with T.expect_download() as dl:T.click("#acTpl")
  fp=SP+"/tpl_empty.xlsx";dl.value.save_as(fp);print("file:",dl.value.suggested_filename)
  wb=openpyxl.load_workbook(fp);ws=wb["학사일정"];print("sheets:",wb.sheetnames)
  for r in ws.iter_rows(min_row=1,max_row=8,values_only=True):print("  ",r)
  print("validation:",[ (str(dv.sqref)[:20],dv.formula1) for dv in ws.data_validations.dataValidation][:2])
  # teacher fills: delete examples, add rows (date cells and text)
  for r in range(5,12):
    for cc in range(1,11):ws.cell(r,cc).value=None
  data=[(datetime.datetime(2026,11,13),None,"개교기념일","수업 없음",None,""),(datetime.datetime(2026,12,23),datetime.datetime(2026,12,24),"기말 이후 단축","단축수업",35,""),("2026.11.26",None,"교내 과학탐구대회","",None,""),("12/31","","겨울방학식","",None,""),(datetime.datetime(2026,11,30),None,"교직원 연수","일정만",None,""),(datetime.datetime(2026,11,6),None,"메모만","넣지 않음",None,"")]
  for i,row in enumerate(data):
    for j,v in enumerate(row):ws.cell(5+i,[1,2,3,4,9,10][j]).value=v
  wb.save(SP+"/tpl_filled.xlsx")
  T.set_input_files("#acIn",SP+"/tpl_filled.xlsx");T.wait_for_timeout(800)
  print("info:",T.inner_text(".cginfo") if T.is_visible(".cginfo") else "-")
  rows=T.evaluate("[...document.querySelectorAll('#acList .acrow')].map(r=>r.innerText.replace(/\\n/g,' ')+' | '+r.querySelector('select').selectedOptions[0].text)");[print("   ",x) for x in rows]
  T.screenshot(path=SP+"/shot/ac_tpl_review.png")
  T.click("#acOk");T.wait_for_timeout(300);print("toast:",T.evaluate("(document.querySelector('.toast')||{}).textContent"))
  st=T.evaluate("JSON.parse(localStorage.getItem('att-app-v2'))");t=st['tt']['list'][0]
  print("chg:",[(x['from'],x['to'],x['reason']) for x in t['chg'] if x['scope']=='day'],"days:",{k:(v['classMin'],v['note']) for k,v in t['days'].items()},"events:",[(e['date'],e['title']) for e in st['events'] if e['id'].startswith('ac')])
  # download filled template (round trip)
  T.click("#acUp2")
  with T.expect_download() as dl:T.click("#acTpl")
  fp=SP+"/tpl_round.xlsx";dl.value.save_as(fp);print("file2:",dl.value.suggested_filename)
  ws=openpyxl.load_workbook(fp)["학사일정"]
  for r in ws.iter_rows(min_row=5,max_row=10,values_only=True):print("  ",r)
  # re-upload round file → no duplicates
  T.set_input_files("#acIn",fp);T.wait_for_timeout(800);T.check("#acPast") if T.is_visible("#acPast") else None;T.click("#acOk");T.wait_for_timeout(300)
  st=T.evaluate("JSON.parse(localStorage.getItem('att-app-v2'))");t=st['tt']['list'][0]
  print("after round: chg",len([x for x in t['chg'] if x['scope']=='day']),"days",len(t['days']),"events",len([e for e in st['events'] if e['id'].startswith('ac')]))
  b.close()
print(errs)
